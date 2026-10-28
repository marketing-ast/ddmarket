from __future__ import annotations

import math
import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from import_merge import (
    calculate_safe_quantity,
    clean_text,
    convert_incoming_price_for_unit,
    extract_consolidated_records,
    format_number,
    is_valid_barcode,
    is_weighted_unit,
    name_similarity,
    normalize_barcode,
    parse_decimal,
    parse_price,
)


PRODUCTS_SHEET = "products"
REPORT_SHEET = "import_report"
DEFAULT_SPREADSHEET_ID = "1IcNPMvcEHIkpIIX9V1Ll7iax3zwVfnORNph83-haaWM"
REQUIRED_PRODUCT_COLUMNS = ("barcode", "name1s", "price", "ostatok", "unit")
INVENTORY_COLUMNS = ("ostatok_1c", "reserve", "min_qty", "ruchnoi_ostatok")
LEGACY_COLUMN_ALIASES = {
    "quantity": "ostatok",
    "quantity_1c": "ostatok_1c",
    "manual_quantity": "ruchnoi_ostatok",
}
NAME_SIMILARITY_THRESHOLD = 0.70
MIN_MATCH_RATIO_FOR_ZEROING = 0.50
PRICE_CHANGE_REVIEW_THRESHOLD = Decimal("0.20")
MAX_RUCHNOI_OSTATOK = Decimal("1000")
PRICE_REVIEW_SHEET = "Изменения цены >20%"


class SheetConflictError(RuntimeError):
    """The sheet changed after the user reviewed the proposed operation."""


def product_snapshot_fingerprint(values: list[list[Any]]) -> str:
    headers = _headers(values)
    relevant = ("id", "publish", "barcode", "name", "name1s", "unit", "price",
                "ostatok_1c", "reserve", "min_qty", "ruchnoi_ostatok")
    indexes = [headers.index(name) if name in headers else None for name in relevant]
    rows = [[clean_text(row[i]) if i is not None and i < len(row) else "" for i in indexes]
            for row in values[1:]]
    while rows and not any(rows[-1]):
        rows.pop()
    # Ignore derived formulas/photos/category formatting; retain row order and identity.
    encoded = json.dumps([headers, rows], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode()).hexdigest()


@dataclass
class SheetsConfig:
    spreadsheet_id: str
    credentials_path: str


def load_sheets_config() -> SheetsConfig | None:
    credentials_path = (
        os.environ.get("DDMARKET_GOOGLE_SERVICE_ACCOUNT_JSON")
        or os.environ.get("GOOGLE_APPLICATION_CREDENTIALS")
        or ""
    )
    spreadsheet_id = os.environ.get("DDMARKET_SPREADSHEET_ID") or DEFAULT_SPREADSHEET_ID
    if not credentials_path or not Path(credentials_path).exists():
        return None
    return SheetsConfig(spreadsheet_id=spreadsheet_id, credentials_path=credentials_path)


def _headers(values: list[list[Any]]) -> list[str]:
    if not values:
        return []
    return [LEGACY_COLUMN_ALIASES.get(clean_text(value).lower(), clean_text(value).lower()) for value in values[0]]


def _row_as_dict(headers: list[str], row: list[Any], row_number: int) -> dict[str, Any]:
    result = {header: row[index] if index < len(row) else "" for index, header in enumerate(headers)}
    result["_row_number"] = row_number
    return result


def _report_issue(report: dict[str, Any], bucket: str, payload: dict[str, Any]) -> None:
    report["issues"].setdefault(bucket, []).append(payload)


def parse_ruchnoi_ostatok(value: object, unit: object) -> Decimal | None | str:
    if clean_text(value) == "":
        return None
    quantity = parse_decimal(value)
    if quantity is None or quantity < 0:
        return "ruchnoi_ostatok: укажите число не меньше 0 или оставьте поле пустым"
    if is_weighted_unit(unit):
        if quantity.quantize(Decimal("0.1")) != quantity:
            return "ruchnoi_ostatok для весового товара: число с точностью до 0,1"
    elif quantity != quantity.to_integral_value():
        return "ruchnoi_ostatok для штучного товара: целое число"
    return quantity


def validate_ruchnoi_ostatok(value: object, unit: object) -> Decimal | str:
    quantity = parse_ruchnoi_ostatok(value, unit)
    if quantity is None:
        return "ruchnoi_ostatok: укажите число от 0 до 1000"
    if isinstance(quantity, str):
        return quantity
    if quantity > MAX_RUCHNOI_OSTATOK:
        return "ruchnoi_ostatok: допустимо значение от 0 до 1000"
    return quantity


def ostatok_formula(headers: list[str], row_number: int) -> str:
    indexes = {header: index + 1 for index, header in enumerate(headers)}
    required = ("unit", "ostatok_1c", "reserve", "min_qty")
    missing = [column for column in required if column not in indexes]
    if missing:
        raise RuntimeError(f"В products нет колонок для формулы ostatok: {', '.join(missing)}")
    unit = f"{column_name(indexes['unit'])}{row_number}"
    ostatok_1c = f"{column_name(indexes['ostatok_1c'])}{row_number}"
    reserve = f"{column_name(indexes['reserve'])}{row_number}"
    min_qty = f"{column_name(indexes['min_qty'])}{row_number}"
    weighted = (
        f'OR(LOWER({unit})="кг";LOWER({unit})="kg";LOWER({unit})="0,1";'
        f'LOWER({unit})="0.1";LOWER({unit})="100 г";LOWER({unit})="100г";LOWER({unit})="100g")'
    )
    return (
        f'=IF({ostatok_1c}="";0;'
        f'LET(raw;{ostatok_1c};reserve;IF({reserve}="";0;{reserve})/100;weighted;{weighted};'
        'safe;IF(weighted;ROUNDDOWN(MAX(0;raw*(1-reserve));1);ROUNDDOWN(MAX(0;raw*(1-reserve));0));'
        f'IF(AND({min_qty}<>"";safe<{min_qty});0;safe)))'
    )


def build_update_report(
    sheet_values: list[list[Any]],
    source_rows: list[dict[str, Any]],
    *,
    allow_large_price_changes: bool = False,
) -> tuple[list[str], dict[str, Any]]:
    if not sheet_values:
        raise RuntimeError("Лист products пуст.")

    headers = _headers(sheet_values)
    missing = [column for column in REQUIRED_PRODUCT_COLUMNS if column not in headers]
    if missing:
        raise RuntimeError(f"В Google Sheets нет обязательных колонок: {', '.join(missing)}.")

    sheet_rows = [_row_as_dict(headers, row, row_number) for row_number, row in enumerate(sheet_values[1:], start=2)]
    product_rows = [row for row in sheet_rows if clean_text(row.get("id")) or clean_text(row.get("name"))]
    master_rows = [
        row
        for row in product_rows
        if clean_text(row.get("publish")).lower() == "yes" and is_valid_barcode(row.get("barcode"))
    ]
    master_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in master_rows:
        master_groups[normalize_barcode(row.get("barcode"))].append(row)

    report: dict[str, Any] = {
        "fatal_errors": [],
        "updates": [],
        "issues": {},
        "summary": {
            "google_product_rows": len(product_rows),
            "google_master_rows": len(master_rows),
            "google_master_without_barcode": sum(
                clean_text(row.get("publish")).lower() == "yes" and not is_valid_barcode(row.get("barcode"))
                for row in product_rows
            ),
            "found_in_1c": 0,
            "ready_to_update": 0,
            "not_found_in_1c": 0,
            "zeroed": 0,
            "ruchnoi_ostatok_rows": 0,
            "price_review_required": 0,
            "price_review_applied": 0,
            "allow_large_price_changes": allow_large_price_changes,
        },
    }

    source_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for source_row in source_rows:
        barcode = normalize_barcode(source_row.get("barcode"))
        if is_valid_barcode(barcode):
            source_groups[barcode].append(source_row)

    blocked_google_barcodes: set[str] = set()
    for barcode, rows in master_groups.items():
        if len(rows) <= 1:
            continue
        blocked_google_barcodes.add(barcode)
        _report_issue(
            report,
            "google_duplicate_barcodes",
            {
                "barcode": barcode,
                "google_rows": ", ".join(str(row["_row_number"]) for row in rows),
                "details": "несколько строк с одним штрихкодом; обновление пропущено",
            },
        )

    syncable_barcodes: set[str] = set()
    matched_barcodes: set[str] = set()

    # Google Sheet is the master: every decision starts with one existing Google row.
    # Extra products from the 1C export are intentionally ignored.
    for barcode, rows in master_groups.items():
        google = rows[0]
        if barcode in blocked_google_barcodes:
            continue
        syncable_barcodes.add(barcode)
        ruchnoi_ostatok = parse_ruchnoi_ostatok(google.get("ruchnoi_ostatok"), google.get("unit"))
        if isinstance(ruchnoi_ostatok, str):
            _report_issue(
                report,
                "invalid_ruchnoi_ostatok",
                {
                    "google_row": google["_row_number"],
                    "barcode": barcode,
                    "details": f"{ruchnoi_ostatok}; будет использован автоматический остаток из 1С",
                },
            )
            ruchnoi_ostatok = None
        has_ruchnoi_ostatok = ruchnoi_ostatok is not None
        if has_ruchnoi_ostatok:
            report["summary"]["ruchnoi_ostatok_rows"] += 1

        candidates = source_groups.get(barcode, [])
        if not candidates:
            report["summary"]["not_found_in_1c"] += 1
            report["summary"]["zeroed"] += 1
            report["updates"].append(
                {
                    "row_number": google["_row_number"],
                    "barcode": barcode,
                    "name": clean_text(google.get("name")),
                    "price": None,
                    "ostatok_1c": 0,
                    "ostatok": 0,
                    "reason": "absent_from_1c_ruchnoi_ostatok" if has_ruchnoi_ostatok else "absent_from_1c",
                }
            )
            _report_issue(
                report,
                "not_found_in_1c",
                {
                    "google_row": google["_row_number"],
                    "barcode": barcode,
                    "name": clean_text(google.get("name")),
                    "details": "в файле 1С строки нет; ostatok_1c и ostatok будут установлены в 0"
                    + (", но сайт продолжит использовать ручной остаток" if has_ruchnoi_ostatok else ""),
                },
            )
            continue
        if len(candidates) > 1:
            _report_issue(
                report,
                "source_duplicate_barcodes",
                {
                    "google_row": google["_row_number"],
                    "barcode": barcode,
                    "source_rows": ", ".join(str(row["source_row"]) for row in candidates),
                    "details": "в файле 1С несколько строк; данные не применяются",
                },
            )
            continue

        source = candidates[0]
        source_row_number = source["source_row"]
        report["summary"]["found_in_1c"] += 1
        matched_barcodes.add(barcode)
        if source["retail_price"] is None or source["retail_price"] <= 0:
            _report_issue(report, "invalid_price", {"source_row": source_row_number, "barcode": barcode, "details": "не заполнена розничная цена"})
            continue
        google_name1s = clean_text(google.get("name1s"))
        similarity = name_similarity(source["name1s"], google_name1s) if google_name1s else NAME_SIMILARITY_THRESHOLD
        if similarity < NAME_SIMILARITY_THRESHOLD:
            _report_issue(
                report,
                "name_similarity_low",
                {
                    "source_row": source_row_number,
                    "barcode": barcode,
                    "source_name1s": source["name1s"],
                    "google_name1s": clean_text(google.get("name1s")),
                    "similarity": f"{similarity:.0%}",
                    "details": "совпадение name1s ниже 70%; строка не применяется",
                },
            )
            continue

        incoming_price = convert_incoming_price_for_unit(
            int(source["retail_price"].quantize(Decimal("1"), rounding=ROUND_HALF_UP)),
            google.get("unit"),
        )
        current_price = parse_decimal(google.get("price"))
        price_change_ratio = (
            abs(Decimal(incoming_price) - current_price) / current_price
            if current_price is not None and current_price > 0
            else None
        )
        if source["main_quantity"] is None or source["total_quantity"] is None:
            _report_issue(report, "empty_quantity", {"source_row": source_row_number, "barcode": barcode, "details": "один или оба остатка пустые"})
            continue
        if source["main_quantity"] != source["total_quantity"]:
            _report_issue(
                report,
                "quantity_mismatch",
                {
                    "source_row": source_row_number,
                    "barcode": barcode,
                    "main_quantity": format_number(source["main_quantity"]),
                    "total_quantity": format_number(source["total_quantity"]),
                    "details": "остатки «Основной» и «Итого» различаются; строка не применяется",
                },
            )
            continue
        calculated = calculate_safe_quantity(
            source["total_quantity"],
            google.get("reserve"),
            google.get("min_qty"),
            google.get("unit"),
        )
        if isinstance(calculated, str):
            _report_issue(report, "invalid_correction", {"source_row": source_row_number, "barcode": barcode, "details": calculated})
            continue
        raw_quantity, final_quantity = calculated
        if raw_quantity < 0:
            _report_issue(
                report,
                "negative_quantity",
                {"source_row": source_row_number, "barcode": barcode, "details": f"из 1С пришло {format_number(raw_quantity)}; на сайт будет записан 0"},
            )
        requires_price_confirmation = (
            price_change_ratio is not None and price_change_ratio > PRICE_CHANGE_REVIEW_THRESHOLD
        )
        if requires_price_confirmation:
            report["summary"]["price_review_required"] += 1
            _report_issue(
                report,
                "price_change_over_20_percent",
                {
                    "google_row": google["_row_number"],
                    "source_row": source_row_number,
                    "barcode": barcode,
                    "name": clean_text(google.get("name")) or google_name1s or clean_text(source.get("name1s")),
                    "current_price": format_number(current_price),
                    "incoming_price": format_number(incoming_price),
                    "change": f"{price_change_ratio:.0%}",
                    "details": (
                        "цена отличается от текущей более чем на 20%; "
                        + (
                            "применена по явному подтверждению"
                            if allow_large_price_changes
                            else "требует отдельного подтверждения"
                        )
                    ),
                },
            )
            if not allow_large_price_changes:
                continue
            report["summary"]["price_review_applied"] += 1

        reason = "ruchnoi_ostatok_auto_calculated" if has_ruchnoi_ostatok else "matched"
        if requires_price_confirmation:
            reason = "price_change_over_20_percent_confirmed"
        report["updates"].append(
            {
                "row_number": google["_row_number"],
                "barcode": barcode,
                "name": clean_text(google.get("name")),
                "price": incoming_price,
                "ostatok_1c": float(raw_quantity),
                "ostatok": float(final_quantity),
                "reason": reason,
            }
        )
        report["summary"]["ready_to_update"] += 1

    report["summary"]["automatic_master_rows"] = len(syncable_barcodes)
    report["summary"]["match_ratio"] = round(len(matched_barcodes) / len(syncable_barcodes), 4) if syncable_barcodes else 1
    if syncable_barcodes and report["summary"]["match_ratio"] < MIN_MATCH_RATIO_FOR_ZEROING:
        report["fatal_errors"].append(
            "С файлом надёжно сопоставлено слишком мало товаров "
            f"({len(matched_barcodes)} из {len(syncable_barcodes)}). "
            "Импорт отменён, чтобы не обнулить каталог частичным или неверным отчётом."
        )

    report["summary"]["updates"] = len(report["updates"])
    report["summary"]["issues"] = sum(len(rows) for rows in report["issues"].values())
    report["summary"]["can_apply"] = not report["fatal_errors"]
    return headers, report


def format_report_message(report: dict[str, Any]) -> str:
    if report["fatal_errors"]:
        return "Файл не применён:\n" + "\n".join(f"• {error}" for error in report["fatal_errors"])
    summary = report["summary"]
    return "\n".join(
        [
            "Проверка каталога Google Sheet по консолидированному отчёту 1С завершена.",
            f"Товарных строк в Google Sheet: {summary['google_product_rows']}",
            f"Мастер-строк для синхронизации: {summary['google_master_rows']}",
            f"Без штрихкода среди publish=yes: {summary['google_master_without_barcode']}",
            f"Ручной остаток задан: {summary['ruchnoi_ostatok_rows']}",
            f"Строк для расчёта по 1С: {summary['automatic_master_rows']}",
            f"Данные найдены в 1С: {summary['found_in_1c']}",
            f"Готово к обновлению: {summary['ready_to_update']}",
            f"Не найдено в 1С: {summary['not_found_in_1c']}",
            f"Будет обнулено: {summary['zeroed']}",
            f"Цен больше 20%, требующих подтверждения: {summary['price_review_required']}",
            f"Цен больше 20%, разрешённых явно: {summary['price_review_applied']}",
            f"Покрытие строк данными 1С: {summary['match_ratio']:.0%}",
            f"Замечаний в полном отчёте: {summary['issues']}",
        ]
    )


def _issue_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for category, issues in report["issues"].items():
        for issue in issues:
            rows.append({"category": category, **issue})
    return rows


def _excel_number(value: object) -> int | float | str:
    number = parse_decimal(value)
    if number is None:
        return clean_text(value)
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _price_review_rows(report: dict[str, Any]) -> list[dict[str, Any]]:
    return list(report.get("issues", {}).get("price_change_over_20_percent", []))


def write_report_workbook(report: dict[str, Any], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "summary"
    summary_sheet.append(["metric", "value"])
    for key, value in report["summary"].items():
        summary_sheet.append([key, value])
    for error in report["fatal_errors"]:
        summary_sheet.append(["fatal_error", error])

    updates = workbook.create_sheet("updates")
    update_headers = ["row_number", "barcode", "name", "price", "ostatok_1c", "ostatok", "reason"]
    updates.append(update_headers)
    for update in report["updates"]:
        updates.append([update.get(key, "") for key in update_headers])

    issues_sheet = workbook.create_sheet("issues")
    issue_rows = _issue_rows(report)
    fields = sorted({key for row in issue_rows for key in row}) or ["category", "details"]
    issues_sheet.append(fields)
    for row in issue_rows:
        issues_sheet.append([row.get(field, "") for field in fields])

    price_review_sheet = workbook.create_sheet(PRICE_REVIEW_SHEET)
    price_review_sheet.append(
        [
            "Строка Google",
            "Строка 1С",
            "Штрихкод",
            "Наименование",
            "Текущая цена",
            "Новая цена",
            "Изменение",
            "Статус",
        ]
    )
    for row in _price_review_rows(report):
        change = parse_decimal(row.get("change"))
        price_review_sheet.append(
            [
                row.get("google_row", ""),
                row.get("source_row", ""),
                row.get("barcode", ""),
                row.get("name", ""),
                _excel_number(row.get("current_price")),
                _excel_number(row.get("incoming_price")),
                float(change / Decimal("100")) if change is not None else "",
                row.get("details", ""),
            ]
        )

    price_review_sheet.freeze_panes = "A2"
    price_review_sheet.auto_filter.ref = price_review_sheet.dimensions
    price_review_sheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in price_review_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in price_review_sheet["C"][1:]:
        cell.number_format = "@"
    for cell in price_review_sheet["E"][1:] + price_review_sheet["F"][1:]:
        cell.number_format = "#,##0"
    for cell in price_review_sheet["G"][1:]:
        cell.number_format = "0%"

    for sheet in workbook.worksheets:
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 70)
    price_review_sheet.column_dimensions["C"].width = max(
        price_review_sheet.column_dimensions["C"].width or 0, 18
    )
    workbook.save(output_path)
    return output_path


def report_to_import_report_values(report: dict[str, Any], applied: bool = False) -> list[list[Any]]:
    values = [["section", "field", "value", "barcode", "details"]]
    for key, value in report["summary"].items():
        values.append(["summary", key, value, "", ""])
    values.append(["summary", "applied", "yes" if applied else "no", "", ""])
    for error in report["fatal_errors"]:
        values.append(["fatal_error", "error", "", "", error])
    for row in _issue_rows(report):
        values.append([row.get("category", ""), row.get("source_row", row.get("google_row", "")), "", row.get("barcode", ""), row.get("details", "")])
    return values


class GoogleSheetsClient:
    def __init__(self, config: SheetsConfig):
        from google.oauth2 import service_account
        from googleapiclient.discovery import build

        credentials = service_account.Credentials.from_service_account_file(
            config.credentials_path,
            scopes=["https://www.googleapis.com/auth/spreadsheets"],
        )
        self.spreadsheet_id = config.spreadsheet_id
        self.service = build("sheets", "v4", credentials=credentials, cache_discovery=False)

    def get_product_values(self) -> list[list[Any]]:
        result = self.service.spreadsheets().values().get(
            spreadsheetId=self.spreadsheet_id,
            range=f"{PRODUCTS_SHEET}!A:Z",
            valueRenderOption="UNFORMATTED_VALUE",
        ).execute()
        return result.get("values", [])

    def find_manual_stock_candidates(self, barcode_suffix: str) -> list[dict[str, Any]]:
        suffix = normalize_barcode(barcode_suffix)
        if not re.fullmatch(r"\d{7,20}", suffix):
            raise ValueError("Введите один штрихкод: только цифры, минимум 7 знаков.")

        values = self.get_product_values()
        if not values:
            raise RuntimeError("Лист products пуст.")
        headers = _headers(values)
        required = ("id", "publish", "barcode", "name", "unit", "ruchnoi_ostatok")
        missing = [column for column in required if column not in headers]
        if missing:
            raise RuntimeError(f"В products нет колонок: {', '.join(missing)}")

        candidates: list[dict[str, Any]] = []
        for row_number, row in enumerate(values[1:], start=2):
            product = _row_as_dict(headers, row, row_number)
            barcode = normalize_barcode(product.get("barcode"))
            if clean_text(product.get("publish")).lower() != "yes":
                continue
            if not is_valid_barcode(barcode) or not barcode.endswith(suffix):
                continue
            candidates.append(
                {
                    "row_number": row_number,
                    "id": clean_text(product.get("id")),
                    "barcode": barcode,
                    "name": clean_text(product.get("name")),
                    "unit": clean_text(product.get("unit")),
                    "ruchnoi_ostatok": clean_text(product.get("ruchnoi_ostatok")),
                }
            )
        return candidates

    def set_ruchnoi_ostatok(
        self,
        row_number: int,
        product_id: str,
        barcode: str,
        value: object,
        *,
        expected_value: object,
        expected_unit: object,
    ) -> dict[str, Any]:
        values = self.get_product_values()
        if not values:
            raise RuntimeError("Лист products пуст.")
        headers = _headers(values)
        required = ("id", "publish", "barcode", "name", "unit", "ruchnoi_ostatok")
        missing = [column for column in required if column not in headers]
        if missing:
            raise RuntimeError(f"В products нет колонок: {', '.join(missing)}")
        if row_number < 2 or row_number > len(values):
            raise RuntimeError("Товар в таблице не найден. Выбери его заново.")

        product = _row_as_dict(headers, values[row_number - 1], row_number)
        if (
            clean_text(product.get("id")) != clean_text(product_id)
            or normalize_barcode(product.get("barcode")) != normalize_barcode(barcode)
            or clean_text(product.get("publish")).lower() != "yes"
        ):
            raise SheetConflictError("Таблица изменилась. Выбери товар заново.")

        if (parse_ruchnoi_ostatok(product.get("ruchnoi_ostatok"), product.get("unit"))
                != parse_ruchnoi_ostatok(expected_value, expected_unit)
                or clean_text(product.get("unit")) != clean_text(expected_unit)):
            raise SheetConflictError("Остаток или единицу измерения уже изменили. Ничего не записал; выбери товар заново и проверь текущее значение.")

        quantity = validate_ruchnoi_ostatok(value, product.get("unit"))
        if isinstance(quantity, str):
            raise ValueError(quantity)

        manual_column = column_name(headers.index("ruchnoi_ostatok") + 1)
        # Re-check identity and the reviewed cells immediately before the write.
        latest = self.get_product_values()
        if _headers(latest) != headers or row_number > len(latest):
            raise SheetConflictError("Таблица изменилась во время подтверждения. Выбери товар заново.")
        current = _row_as_dict(headers, latest[row_number - 1], row_number)
        if any(clean_text(current.get(k)) != clean_text(product.get(k))
               for k in ("id", "barcode", "publish", "unit", "ruchnoi_ostatok")):
            raise SheetConflictError("Товар изменился во время подтверждения. Ничего не записал; выбери его заново.")
        self.service.spreadsheets().values().update(
            spreadsheetId=self.spreadsheet_id,
            range=f"{PRODUCTS_SHEET}!{manual_column}{row_number}",
            valueInputOption="USER_ENTERED",
            body={"values": [[float(quantity)]]},
        ).execute()
        verified = self.get_product_values()
        actual = _row_as_dict(_headers(verified), verified[row_number - 1], row_number) if row_number <= len(verified) else {}
        if (clean_text(actual.get("id")) != clean_text(product_id)
                or normalize_barcode(actual.get("barcode")) != normalize_barcode(barcode)
                or parse_ruchnoi_ostatok(actual.get("ruchnoi_ostatok"), actual.get("unit")) != quantity):
            raise SheetConflictError("После записи данные снова изменились. Проверь текущий остаток; автоматический повтор не выполнялся.")
        return {
            "row_number": row_number,
            "id": clean_text(product.get("id")),
            "barcode": normalize_barcode(product.get("barcode")),
            "name": clean_text(product.get("name")),
            "unit": clean_text(product.get("unit")),
            "ruchnoi_ostatok": format_number(quantity),
        }

    def ensure_inventory_columns(self) -> list[list[Any]]:
        values = self.get_product_values()
        if not values:
            raise RuntimeError("Лист products пуст.")
        headers = _headers(values)
        missing = [column for column in INVENTORY_COLUMNS if column not in headers]
        if not missing:
            return values
        required_column_count = len(headers) + len(missing)
        spreadsheet = self.service.spreadsheets().get(
            spreadsheetId=self.spreadsheet_id,
            fields="sheets(properties(sheetId,title,gridProperties(columnCount)))",
        ).execute()
        properties = next(
            (
                sheet.get("properties", {})
                for sheet in spreadsheet.get("sheets", [])
                if sheet.get("properties", {}).get("title") == PRODUCTS_SHEET
            ),
            None,
        )
        if properties is None:
            raise RuntimeError(f"Не найден лист {PRODUCTS_SHEET}.")
        current_column_count = properties.get("gridProperties", {}).get("columnCount", 0)
        if current_column_count < required_column_count:
            self.service.spreadsheets().batchUpdate(
                spreadsheetId=self.spreadsheet_id,
                body={
                    "requests": [
                        {
                            "updateSheetProperties": {
                                "properties": {
                                    "sheetId": properties["sheetId"],
                                    "gridProperties": {"columnCount": required_column_count},
                                },
                                "fields": "gridProperties.columnCount",
                            }
                        }
                    ]
                },
            ).execute()
        start_column = column_name(len(headers) + 1)
        end_column = column_name(len(headers) + len(missing))
        self.service.spreadsheets().values().update(
            spreadsheetId=self.spreadsheet_id,
            range=f"{PRODUCTS_SHEET}!{start_column}1:{end_column}1",
            valueInputOption="USER_ENTERED",
            body={"values": [missing]},
        ).execute()
        return self.get_product_values()

    def set_reserve_dropdown_and_default(self) -> int:
        values = self.ensure_inventory_columns()
        headers = _headers(values)
        if "reserve" not in headers:
            raise RuntimeError("В products нет колонки reserve.")
        spreadsheet = self.service.spreadsheets().get(
            spreadsheetId=self.spreadsheet_id,
            fields="sheets(properties(sheetId,title,gridProperties(rowCount)))",
        ).execute()
        properties = next(
            (
                sheet.get("properties", {})
                for sheet in spreadsheet.get("sheets", [])
                if sheet.get("properties", {}).get("title") == PRODUCTS_SHEET
            ),
            None,
        )
        if properties is None:
            raise RuntimeError(f"Не найден лист {PRODUCTS_SHEET}.")
        reserve_column_index = headers.index("reserve")
        self.service.spreadsheets().batchUpdate(
            spreadsheetId=self.spreadsheet_id,
            body={
                "requests": [
                    {
                        "setDataValidation": {
                            "range": {
                                "sheetId": properties["sheetId"],
                                "startRowIndex": 1,
                                "endRowIndex": properties.get("gridProperties", {}).get("rowCount", len(values)),
                                "startColumnIndex": reserve_column_index,
                                "endColumnIndex": reserve_column_index + 1,
                            },
                            "rule": {
                                "condition": {
                                    "type": "ONE_OF_LIST",
                                    "values": [{"userEnteredValue": value} for value in ("0", "5", "10", "15", "20")],
                                },
                                "strict": True,
                                "showCustomUi": True,
                                "inputMessage": "Резерв: 0, 5, 10, 15 или 20%",
                            },
                        }
                    }
                ]
            },
        ).execute()
        product_row_count = max(len(values) - 1, 0)
        if product_row_count:
            reserve_column = column_name(reserve_column_index + 1)
            self.service.spreadsheets().values().update(
                spreadsheetId=self.spreadsheet_id,
                range=f"{PRODUCTS_SHEET}!{reserve_column}2:{reserve_column}{product_row_count + 1}",
                valueInputOption="USER_ENTERED",
                body={"values": [[10] for _ in range(product_row_count)]},
            ).execute()
        return product_row_count

    def ensure_ostatok_formulas(self, headers: list[str], row_count: int) -> None:
        indexes = {header: index + 1 for index, header in enumerate(headers)}
        if "ostatok" not in indexes:
            raise RuntimeError("В products нет колонки ostatok.")
        if row_count <= 1:
            return
        ostatok_column = column_name(indexes["ostatok"])
        self.service.spreadsheets().values().update(
            spreadsheetId=self.spreadsheet_id,
            range=f"{PRODUCTS_SHEET}!{ostatok_column}2:{ostatok_column}{row_count}",
            valueInputOption="USER_ENTERED",
            body={"values": [[ostatok_formula(headers, row_number)] for row_number in range(2, row_count + 1)]},
        ).execute()

    def write_import_report(self, report: dict[str, Any], applied: bool = False) -> None:
        values = report_to_import_report_values(report, applied=applied)
        self.service.spreadsheets().values().clear(
            spreadsheetId=self.spreadsheet_id,
            range=f"{REPORT_SHEET}!A:Z",
            body={},
        ).execute()
        self.service.spreadsheets().values().update(
            spreadsheetId=self.spreadsheet_id,
            range=f"{REPORT_SHEET}!A1",
            valueInputOption="USER_ENTERED",
            body={"values": values},
        ).execute()

    def apply_updates(self, headers: list[str], report: dict[str, Any]) -> int:
        indexes = {header: index + 1 for index, header in enumerate(headers)}
        missing = [column for column in ("price", "ostatok", "ostatok_1c") if column not in indexes]
        if missing:
            raise RuntimeError(f"В products нет колонок: {', '.join(missing)}")
        data = []
        for update in report["updates"]:
            row = update["row_number"]
            if update["price"] is not None:
                data.append({"range": f"{PRODUCTS_SHEET}!{column_name(indexes['price'])}{row}", "values": [[update["price"]]]})
            if update["ostatok_1c"] is not None:
                data.append({"range": f"{PRODUCTS_SHEET}!{column_name(indexes['ostatok_1c'])}{row}", "values": [[update["ostatok_1c"]]]})
        if data:
            self.service.spreadsheets().values().batchUpdate(
                spreadsheetId=self.spreadsheet_id,
                body={"valueInputOption": "USER_ENTERED", "data": data},
            ).execute()
        return len(report["updates"])


def analyze_consolidated_file(path: str | Path, client: GoogleSheetsClient) -> tuple[list[str], dict[str, Any]]:
    values = client.get_product_values()
    headers, report = build_update_report(values, extract_consolidated_records(path))
    report["sheet_fingerprint"] = product_snapshot_fingerprint(values)
    return headers, report


def apply_consolidated_file(
    path: str | Path,
    client: GoogleSheetsClient,
    *,
    allow_large_price_changes: bool = False,
    expected_fingerprint: str | None = None,
) -> tuple[list[str], dict[str, Any], int]:
    sheet_values = client.get_product_values()
    before = product_snapshot_fingerprint(sheet_values)
    if expected_fingerprint is not None and before != expected_fingerprint:
        raise SheetConflictError("Таблица изменилась после проверки. Импорт не выполнен; нажми «Проверить файл» и подтверди новый отчёт.")
    if any(column not in _headers(sheet_values) for column in INVENTORY_COLUMNS):
        if expected_fingerprint is not None:
            raise SheetConflictError("Изменилась структура таблицы. Импорт не выполнен; требуется проверка колонок.")
        sheet_values = client.ensure_inventory_columns()
        before = product_snapshot_fingerprint(sheet_values)
    headers, report = build_update_report(
        sheet_values,
        extract_consolidated_records(path),
        allow_large_price_changes=allow_large_price_changes,
    )
    if report["fatal_errors"]:
        return headers, report, 0
    if product_snapshot_fingerprint(client.get_product_values()) != before:
        raise SheetConflictError("Таблица изменилась во время подготовки. Импорт не выполнен; проверь файл заново.")
    client.ensure_ostatok_formulas(headers, len(sheet_values))
    if product_snapshot_fingerprint(client.get_product_values()) != before:
        raise SheetConflictError("Таблица изменилась перед записью. Цены и остатки не записаны; проверь файл заново.")
    applied_rows = client.apply_updates(headers, report)
    return headers, report, applied_rows


def column_name(index: int) -> str:
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name
