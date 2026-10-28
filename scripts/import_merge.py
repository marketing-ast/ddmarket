from __future__ import annotations

import math
import re
from decimal import Decimal, InvalidOperation, ROUND_DOWN, ROUND_HALF_UP
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ALLOWED_RESERVES = {Decimal(value) for value in ("0", "5", "10", "15", "20")}


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def normalize_header(value: object) -> str:
    return clean_text(value).lower().replace("ё", "е")


def normalize_barcode(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value).lstrip("'").replace(" ", "")


def is_valid_barcode(value: object) -> bool:
    return bool(re.fullmatch(r"\d{3,20}", normalize_barcode(value)))


def parse_price(value: object) -> int | None:
    number = parse_decimal(value)
    if number is None or number <= 0:
        return None
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_decimal(value: object) -> Decimal | None:
    if value is None or clean_text(value) == "":
        return None
    try:
        normalized = clean_text(value).replace(" ", "").replace(",", ".").replace("%", "")
        return Decimal(normalized)
    except (InvalidOperation, ValueError):
        return None


def format_number(value: Decimal | float | int) -> str:
    decimal = Decimal(str(value))
    rendered = format(decimal.normalize(), "f")
    return "0" if rendered in {"", "-0"} else rendered


def normalize_name(value: object) -> str:
    normalized = clean_text(value).lower()
    return re.sub(r"\s*,\s*(шт|кг)\s*$", "", normalized).strip()


def name_similarity(left: object, right: object) -> float:
    normalized_left = normalize_name(left)
    normalized_right = normalize_name(right)
    if not normalized_left or not normalized_right:
        return 0.0
    return SequenceMatcher(None, normalized_left, normalized_right).ratio()


def is_100_gram_unit(value: object) -> bool:
    normalized = clean_text(value).lower().replace(",", ".")
    return normalized in {"0.1", "100 г", "100г", "100g"}


def is_weighted_unit(value: object) -> bool:
    normalized = clean_text(value).lower().replace(",", ".")
    return is_100_gram_unit(normalized) or "кг" in normalized or "kg" in normalized


def convert_incoming_price_for_unit(price: int, unit: object) -> int:
    return math.ceil(price / 10) if is_100_gram_unit(unit) else price


def calculate_safe_quantity(
    raw_quantity: Decimal,
    reserve_value: object,
    min_qty_value: object,
    unit: object,
) -> tuple[Decimal, Decimal] | str:
    reserve = parse_decimal(reserve_value)
    if reserve is None:
        reserve = Decimal("0")
    if reserve not in ALLOWED_RESERVES:
        return "reserve: допустимы только 0, 5, 10, 15 или 20"

    weighted = is_weighted_unit(unit)
    min_qty = parse_decimal(min_qty_value)
    if min_qty is not None:
        if weighted:
            if min_qty <= 0 or min_qty > 10 or min_qty.quantize(Decimal("0.1")) != min_qty:
                return "min_qty для весового товара: от 0,1 до 10 с шагом 0,1"
        elif min_qty < 1 or min_qty > 10 or min_qty != min_qty.to_integral_value():
            return "min_qty для штучного товара: целое от 1 до 10"

    if raw_quantity <= 0:
        return raw_quantity, Decimal("0")

    safe_quantity = raw_quantity * (Decimal("1") - reserve / Decimal("100"))
    if weighted:
        safe_quantity = safe_quantity.quantize(Decimal("0.1"), rounding=ROUND_DOWN)
    else:
        safe_quantity = safe_quantity.to_integral_value(rounding=ROUND_DOWN)
    if min_qty is not None and safe_quantity < min_qty:
        safe_quantity = Decimal("0")
    return raw_quantity, safe_quantity


def extract_consolidated_records(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        profile = find_consolidated_profile(rows)
        if profile is None:
            raise ValueError(
                "Не найдены колонки 1С: «Номенклатура, Ед. изм.», «Штрихкод», «Цена розничная», "
                "«Основной → Количество», «Итого → Количество»."
            )
        header_row, name_column, barcode_column, retail_column, main_quantity_column, total_quantity_column = profile
        records: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
            name1s = clean_text(_cell(row, name_column))
            barcode = normalize_barcode(_cell(row, barcode_column))
            retail_price = parse_decimal(_cell(row, retail_column))
            main_quantity = parse_decimal(_cell(row, main_quantity_column))
            total_quantity = parse_decimal(_cell(row, total_quantity_column))
            if not name1s or (not barcode and retail_price is None):
                continue
            records.append(
                {
                    "source_row": row_number,
                    "barcode": barcode,
                    "name1s": name1s,
                    "retail_price": retail_price,
                    "main_quantity": main_quantity,
                    "total_quantity": total_quantity,
                }
            )
        if not records:
            raise ValueError("В файле не найдено ни одной товарной строки.")
        return records
    finally:
        workbook.close()


def find_consolidated_profile(rows: list[tuple[Any, ...]]) -> tuple[int, int, int, int, int, int] | None:
    for row_number, row in enumerate(rows[:20], start=1):
        headers = [normalize_header(value) for value in row]
        required = {"номенклатура, ед. изм.", "штрихкод", "цена розничная"}
        if not required.issubset(headers):
            continue
        quantity_columns = [index + 1 for index, value in enumerate(headers) if value == "количество"]
        main_column = next(
            (column for column in quantity_columns if row_number > 1 and normalize_header(_cell(rows[row_number - 2], column)) == "основной"),
            None,
        )
        total_column = next(
            (column for column in quantity_columns if row_number > 1 and normalize_header(_cell(rows[row_number - 2], column)) == "итого"),
            None,
        )
        if main_column and total_column:
            return (
                row_number,
                headers.index("номенклатура, ед. изм.") + 1,
                headers.index("штрихкод") + 1,
                headers.index("цена розничная") + 1,
                main_column,
                total_column,
            )
    return None


def _cell(row: tuple[Any, ...], column_number: int) -> Any:
    index = column_number - 1
    return row[index] if index < len(row) else None
